"""ATS Resume Checker - Streamlit app powered by Google Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description, and get
an estimated ATS score with concrete suggestions for improvement.
"""

import io
import json
import os
import re
import time

import streamlit as st
from docx import Document
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "gemini-3.5-flash"  # override with the GEMINI_MODEL secret/env var
FALLBACK_MODEL = "gemini-2.5-flash"  # used if the main model is busy/unavailable
RETRIES_PER_MODEL = 3  # tries per model before moving on
RETRY_WAIT_SECONDS = 2  # doubles after each failed try (2s, 4s)
MAX_FILE_MB = 5
MAX_CHARS = 30_000  # keeps the prompt small and fast
MIN_CHARS = 200  # below this, the file is probably scanned/empty

CATEGORIES = [
    "Keywords & Skills",
    "Formatting & Structure",
    "Work Experience & Impact",
    "Education & Certifications",
    "Contact Info & Summary",
    "Readability & Grammar",
]

PROMPT_TEMPLATE = """You are an expert ATS (Applicant Tracking System) analyst and senior recruiter.
Evaluate the resume below for how well it would parse and rank in an ATS.

{jd_block}

Score each category from 0 to 100 using these categories exactly:
{categories}

Rules:
- Be honest and specific; do not inflate scores.
- Quote or reference actual content from the resume in your suggestions.
- "overall_score" must be an integer from 0 to 100.
- Return ONLY valid JSON, no markdown, matching this schema:

{{
  "overall_score": 0,
  "summary": "2-3 sentence overall assessment",
  "category_scores": {{"<category name>": 0}},
  "strengths": ["..."],
  "missing_keywords": ["..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "area": "...", "issue": "...", "fix": "..."}}
  ],
  "formatting_issues": ["..."],
  "rewritten_examples": [
    {{"original": "weak bullet from the resume", "improved": "stronger rewrite"}}
  ]
}}

RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    ext = os.path.splitext(file_name)[1].lower()

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if ext == ".txt":
        return data.decode("utf-8", errors="ignore")

    raise ValueError(f"Unsupported file type: {ext or 'unknown'}")


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ----------------------------------------------------------------------------
# Gemini call + response handling
# ----------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = job_description.strip()
    if jd:
        jd_block = (
            "A target job description is provided. Weight keyword match and "
            "relevance to it heavily.\n\nJOB DESCRIPTION:\n\"\"\"\n"
            + jd[:8000]
            + "\n\"\"\""
        )
    else:
        jd_block = (
            "No job description was provided. Judge against general ATS best "
            "practices and common expectations for the candidate's apparent field."
        )
    return PROMPT_TEMPLATE.format(
        jd_block=jd_block,
        categories="\n".join(f"- {c}" for c in CATEGORIES),
        resume=resume_text[:MAX_CHARS],
    )


def parse_model_json(raw: str) -> dict:
    """Parse JSON from the model, tolerating code fences or stray text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start : end + 1])
        raise ValueError("The AI response was not valid JSON. Please try again.")


def _clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def normalize_result(data: dict) -> dict:
    """Make sure every field exists and has the expected type."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected AI response format.")

    cat_scores = data.get("category_scores") or {}
    if not isinstance(cat_scores, dict):
        cat_scores = {}

    def str_list(key):
        val = data.get(key) or []
        return [str(v) for v in val if v] if isinstance(val, list) else []

    def dict_list(key):
        val = data.get(key) or []
        return [v for v in val if isinstance(v, dict)] if isinstance(val, list) else []

    return {
        "overall_score": _clamp_score(data.get("overall_score")),
        "summary": str(data.get("summary") or ""),
        "category_scores": {str(k): _clamp_score(v) for k, v in cat_scores.items()},
        "strengths": str_list("strengths"),
        "missing_keywords": str_list("missing_keywords"),
        "improvements": dict_list("improvements"),
        "formatting_issues": str_list("formatting_issues"),
        "rewritten_examples": dict_list("rewritten_examples"),
    }


def _is_busy_error(exc: Exception) -> bool:
    msg = str(exc)
    return any(c in msg for c in ("503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED", "500", "INTERNAL", "504", "DEADLINE"))


def _is_model_missing(exc: Exception) -> bool:
    msg = str(exc)
    return "404" in msg or "NOT_FOUND" in msg


def analyze_resume(
    resume_text: str,
    job_description: str,
    api_key: str,
    model: str,
    fallback_model: str = FALLBACK_MODEL,
) -> dict:
    """Call Gemini; retry when it is busy and fall back to a second model."""
    # Imported here so the rest of the module works without the SDK installed.
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )

    models = [model]
    if fallback_model and fallback_model != model:
        models.append(fallback_model)

    last_exc = None
    for m in models:
        for attempt in range(RETRIES_PER_MODEL):
            try:
                response = client.models.generate_content(
                    model=m, contents=prompt, config=config
                )
                text = getattr(response, "text", None)
                if not text:
                    raise ValueError(
                        "The AI returned an empty response (it may have been blocked). Try again."
                    )
                return normalize_result(parse_model_json(text))
            except Exception as exc:
                last_exc = exc
                if _is_model_missing(exc):
                    break  # this model name doesn't exist -> try the next one
                if not _is_busy_error(exc):
                    raise  # bad key, bad request, etc. -> retrying won't help
                if attempt < RETRIES_PER_MODEL - 1:
                    time.sleep(RETRY_WAIT_SECONDS * (2**attempt))
    raise last_exc


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def get_secret(name: str, default: str = "") -> str:
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass  # no secrets.toml locally
    return os.environ.get(name, default)


def clean_api_key(key: str) -> str:
    """Strip whitespace/quotes that often sneak in when pasting a key."""
    return (key or "").strip().strip("\"'").strip()


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent - likely to pass most ATS filters"
    if score >= 65:
        return "Good - a few fixes will help"
    if score >= 50:
        return "Fair - needs noticeable improvement"
    return "Poor - major changes recommended"


PRIORITY_ICON = {"high": "🔴", "medium": "🟠", "low": "🟢"}


def render_result(result: dict) -> None:
    score = result["overall_score"]

    st.subheader("Your ATS Score")
    col1, col2 = st.columns([1, 3])
    col1.metric("Score", f"{score}/100")
    col2.progress(score / 100)
    col2.caption(score_label(score))
    if result["summary"]:
        st.write(result["summary"])

    if result["category_scores"]:
        st.subheader("Score breakdown")
        for name, val in result["category_scores"].items():
            st.write(f"**{name}** - {val}/100")
            st.progress(val / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["No specific strengths listed."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("🔑 Missing keywords")
        if result["missing_keywords"]:
            st.markdown(", ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.markdown("No major keyword gaps found.")

    st.subheader("🛠️ Improvements")
    order = {"high": 0, "medium": 1, "low": 2}
    items = sorted(
        result["improvements"],
        key=lambda i: order.get(str(i.get("priority", "")).lower(), 3),
    )
    if not items:
        st.markdown("No improvements suggested.")
    for item in items:
        prio = str(item.get("priority", "")).lower()
        icon = PRIORITY_ICON.get(prio, "⚪")
        title = f"{icon} {item.get('priority', '')} - {item.get('area', 'General')}"
        with st.expander(title, expanded=(prio == "high")):
            st.markdown(f"**Issue:** {item.get('issue', '')}")
            st.markdown(f"**Fix:** {item.get('fix', '')}")

    if result["formatting_issues"]:
        st.subheader("📄 Formatting issues")
        for f in result["formatting_issues"]:
            st.markdown(f"- {f}")

    if result["rewritten_examples"]:
        st.subheader("✍️ Example rewrites")
        for ex in result["rewritten_examples"]:
            st.markdown(f"**Before:** {ex.get('original', '')}")
            st.markdown(f"**After:** {ex.get('improved', '')}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# ----------------------------------------------------------------------------
# Main app
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an estimated ATS score and tips to improve it.")

    api_key = clean_api_key(get_secret("GEMINI_API_KEY"))
    model = get_secret("GEMINI_MODEL", DEFAULT_MODEL).strip()

    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        else:
            st.success("API key loaded from secrets.")
        st.markdown("---")
        st.caption(
            "The score is an AI-generated estimate, not the output of a real ATS. "
            "Use it as guidance."
        )

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, but gives a more accurate score)",
        height=150,
        placeholder="Paste the job description here to check keyword match...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = clean_api_key(api_key)
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            return

        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = clean_text(extract_text(uploaded.name, data))
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < MIN_CHARS:
            st.error(
                "Very little text was found. If your resume is a scanned image or "
                "built from images, an ATS can't read it either - export a "
                "text-based PDF or DOCX and try again."
            )
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(
                    text,
                    job_description,
                    api_key,
                    model,
                    get_secret("GEMINI_FALLBACK_MODEL", FALLBACK_MODEL).strip(),
                )
        except Exception as exc:
            msg = str(exc)
            if _is_busy_error(exc):
                st.error(
                    "Google's AI is very busy right now. The app already retried "
                    "and tried a backup model. Please wait a minute and click "
                    "Analyze again."
                )
            elif "401" in msg or "UNAUTHENTICATED" in msg or "API key" in msg:
                st.error(
                    "Google rejected the API key. Create a fresh key at "
                    "https://aistudio.google.com/apikey and paste only the key "
                    "(no quotes or spaces)."
                )
            else:
                st.error(f"Analysis failed: {exc}")
            return

        render_result(result)


if __name__ == "__main__":
    main()
